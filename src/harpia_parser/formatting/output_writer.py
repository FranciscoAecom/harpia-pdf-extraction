from pathlib import Path
from datetime import datetime
import re
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..constants import ACM_EXTRACTION_TIMESTAMP_COLUMN
from ..validation.schemas import VALIDATION_ERROR_COLUMNS, validate_outputs
from .jsonl_writer import write_jsonl
from .excel_writer import write_excel


NUMERIC_TEXT = re.compile(r"^([+-]?\d+(?:[,.]\d+)?)(?:\s*x\s*10\s*([+-]?\d+))?$", re.IGNORECASE)
NUMBER_IN_TEXT = re.compile(r"[<>]?\s*([+-]?\d+(?:[,.]\d+)?)")
SOURCE_NUMBER_IN_TEXT = re.compile(
    r"[<>]?\s*([+-]?\d+(?:[,.]\d+)?)(?:\s*(?:x\s*10|[Ee])\s*[+-]?\d+)?",
    re.IGNORECASE,
)
INTEGER_TEXT = re.compile(r"^\d+$")
DATE_TEXT = re.compile(r"^\d{2}/\d{2}/\d{4}$")
DATE_TIME_TEXT = re.compile(r"^\d{2}/\d{2}/\d{4}\s+\d{2}:\d{2}(?::\d{2})?$")
INTEGER_OUTPUT_COLUMNS = {"id_amostra", "id_taxonomia", "versao_template"}
DATE_OUTPUT_COLUMNS = {
    "results_extract": {"acm_data_inicio"},
    "sample": {"data_coleta", "data_publicacao", "data_recebimento"},
}
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEADER_BORDER = Border(
    left=Side(style="thin", color="9EADCC"),
    right=Side(style="thin", color="9EADCC"),
    top=Side(style="thin", color="9EADCC"),
    bottom=Side(style="thin", color="9EADCC"),
)

def _numeric_value_and_format(value) -> tuple[float, str] | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    match = NUMERIC_TEXT.match(text)
    if not match:
        return None

    mantissa = match.group(1)
    decimal_separator = "," if "," in mantissa else "." if "." in mantissa else None
    decimal_places = len(mantissa.split(decimal_separator, 1)[1]) if decimal_separator else 0
    numeric = float(mantissa.replace(",", "."))
    exponent = match.group(2)
    if exponent is not None:
        numeric *= 10 ** int(exponent)
        number_format = "0" if decimal_places == 0 else f"0.{'0' * decimal_places}"
        return numeric, number_format

    number_format = "0" if decimal_places == 0 else f"0.{'0' * decimal_places}"
    return numeric, number_format


def _number_format_from_text(value) -> str | None:
    parsed = _numeric_value_and_format(value)
    return parsed[1] if parsed else None


def _number_formats_from_source(value) -> list[str]:
    if value is None or pd.isna(value):
        return []
    formats = []
    for match in NUMBER_IN_TEXT.finditer(str(value)):
        number_text = match.group(1)
        number_format = _number_format_from_text(number_text)
        if number_format:
            formats.append(number_format)
    return formats


def _apply_numeric_format_to_column(
    worksheet,
    df: pd.DataFrame,
    column_name: str,
    source_column: str | None = None,
    source_number_index: int = 0,
) -> None:
    if column_name not in df.columns:
        return

    column_index = list(df.columns).index(column_name) + 1
    for row_index, value in enumerate(df[column_name], start=2):
        parsed = _numeric_value_and_format(value)
        if parsed is None and not pd.isna(value):
            try:
                parsed = (float(value), "0")
            except (TypeError, ValueError):
                parsed = None
        if parsed is None:
            continue

        numeric, number_format = parsed
        if source_column and source_column in df.columns:
            source_formats = _number_formats_from_source(df[source_column].iloc[row_index - 2])
            if source_formats:
                number_format = source_formats[min(source_number_index, len(source_formats) - 1)]

        cell = worksheet.cell(row=row_index, column=column_index)
        cell.value = numeric
        cell.number_format = number_format


def _format_numeric_results_sheet(writer: pd.ExcelWriter, df: pd.DataFrame) -> None:
    worksheet = writer.sheets["results_extract"]
    direct_numeric_text_columns = [
        "variacao_percentual",
        "quantidade_adicionada",
        "recuperacao_percentual",
    ]
    for column in direct_numeric_text_columns:
        _apply_numeric_format_to_column(worksheet, df, column)

    source_mapped_columns = [
        ("acm_resultado_tratado", "resultado", 0),
        ("acm_conama_minimo", "conama", 0),
        ("acm_conama_maximo", "conama", 1),
        ("acm_copam_cerh_minimo", "copam_cerh", 0),
        ("acm_copam_cerh_maximo", "copam_cerh", 1),
        ("acm_ld_minimo", "ld", 0),
        ("acm_ld_maximo", "ld", 1),
        ("acm_lq_minimo", "lq", 0),
        ("acm_lq_maximo", "lq", 1),
        ("acm_incerteza_valor", "incerteza", 0),
        ("acm_faixa_aceitacao_minimo", "faixa_aceitacao", 0),
        ("acm_faixa_aceitacao_maximo", "faixa_aceitacao", 1),
    ]
    for column, source, source_number_index in source_mapped_columns:
        _apply_numeric_format_to_column(worksheet, df, column, source, source_number_index)


def _format_output_sheet(writer: pd.ExcelWriter, sheet_name: str) -> None:
    worksheet = writer.sheets.get(sheet_name)
    if worksheet is None or worksheet.max_row < 1:
        return

    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 28
    if worksheet.max_column > 0 and worksheet.max_row > 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = HEADER_BORDER

    _apply_integer_format_to_columns(worksheet)
    _apply_date_format_to_columns(worksheet, sheet_name)

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells[: min(len(column_cells), 200)]:
            value = cell.value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 55)


def _apply_integer_format_to_columns(worksheet) -> None:
    headers = {str(cell.value): cell.column for cell in worksheet[1] if cell.value is not None}
    for column_name in INTEGER_OUTPUT_COLUMNS & set(headers):
        column_index = headers[column_name]
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            if cell.value is None:
                continue
            text = str(cell.value).strip()
            if INTEGER_TEXT.match(text):
                cell.value = int(text)
                cell.number_format = "0"


def _apply_date_format_to_columns(worksheet, sheet_name: str) -> None:
    date_columns = set(DATE_OUTPUT_COLUMNS.get(sheet_name, set()))
    date_columns.add(ACM_EXTRACTION_TIMESTAMP_COLUMN)
    if not date_columns:
        return
    headers = {str(cell.value): cell.column for cell in worksheet[1] if cell.value is not None}
    for column_name in date_columns & set(headers):
        column_index = headers[column_name]
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            if cell.value is None:
                continue
            text = str(cell.value).strip()
            if DATE_TIME_TEXT.match(text):
                date_format = "%d/%m/%Y %H:%M:%S" if text.count(":") == 2 else "%d/%m/%Y %H:%M"
                cell.value = datetime.strptime(text, date_format)
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif DATE_TEXT.match(text):
                cell.value = datetime.strptime(f"{text} 00:00:00", "%d/%m/%Y %H:%M:%S")
                cell.number_format = "dd/mm/yyyy hh:mm"


def _timestamp_text(timestamp: datetime) -> str:
    return timestamp.strftime("%d/%m/%Y %H:%M:%S")


def _with_extraction_timestamp(df: pd.DataFrame | None, timestamp: datetime) -> pd.DataFrame | None:
    if df is None:
        return None
    return _with_required_extraction_timestamp(df, timestamp)


def _with_required_extraction_timestamp(df: pd.DataFrame, timestamp: datetime) -> pd.DataFrame:
    result = df.copy()
    result[ACM_EXTRACTION_TIMESTAMP_COLUMN] = _timestamp_text(timestamp)
    return result


def _write_jsonl_output(
    output_path: Path,
    df: pd.DataFrame,
    sample_df: pd.DataFrame | None,
    client_df: pd.DataFrame | None,
    classification_audit_df: pd.DataFrame | None,
    table_extraction_audit_df: pd.DataFrame | None,
    section_extraction_audit_df: pd.DataFrame | None,
    field_extraction_audit_df: pd.DataFrame | None,
    duplicate_audit_df: pd.DataFrame | None,
    packaging_preservatives_df: pd.DataFrame | None,
    notes_df: pd.DataFrame | None,
    general_considerations_df: pd.DataFrame | None,
    conformity_statement_df: pd.DataFrame | None,
    validation_key_df: pd.DataFrame | None,
    revision_reason_df: pd.DataFrame | None,
    validation_errors: pd.DataFrame,
    sheets_to_write: list[str],
) -> None:
    table_sources: dict[str, pd.DataFrame | None] = {
        "results_extract": df,
        "sample": sample_df,
        "client": client_df,
        "packaging_preservatives": packaging_preservatives_df,
        "notes": notes_df,
        "general_considerations": general_considerations_df,
        "conformity_statement": conformity_statement_df,
        "validation_key": validation_key_df,
        "revision_reason": revision_reason_df,
    }
    audit_sources: dict[str, pd.DataFrame | None] = {
        "classification_audit": classification_audit_df,
        "table_extraction_audit": table_extraction_audit_df,
        "section_extraction_audit": section_extraction_audit_df,
        "field_extraction_audit": field_extraction_audit_df,
        "duplicate_audit": duplicate_audit_df,
        "validation_errors": validation_errors,
    }
    write_jsonl(output_path, table_sources, audit_sources, sheets_to_write)


def salvar(
    df: pd.DataFrame,
    output_path: Path,
    sample_df: pd.DataFrame | None = None,
    client_df: pd.DataFrame | None = None,
    output_tabs: list[str] | None = None,
    classification_audit_df: pd.DataFrame | None = None,
    table_extraction_audit_df: pd.DataFrame | None = None,
    section_extraction_audit_df: pd.DataFrame | None = None,
    field_extraction_audit_df: pd.DataFrame | None = None,
    duplicate_audit_df: pd.DataFrame | None = None,
    packaging_preservatives_df: pd.DataFrame | None = None,
    notes_df: pd.DataFrame | None = None,
    general_considerations_df: pd.DataFrame | None = None,
    conformity_statement_df: pd.DataFrame | None = None,
    validation_key_df: pd.DataFrame | None = None,
    revision_reason_df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    ext = output_path.suffix.lower()
    extraction_timestamp = datetime.now().replace(microsecond=0)
    sheets_to_write = output_tabs or [
        "results_extract",
        "sample",
        "client",
        "packaging_preservatives",
        "notes",
        "general_considerations",
        "conformity_statement",
        "validation_key",
        "revision_reason",
        "table_extraction_audit",
        "section_extraction_audit",
        "field_extraction_audit",
        "classification_audit",
        "duplicate_audit",
        "validation_errors",
    ]
    df = _with_required_extraction_timestamp(df, extraction_timestamp)
    sample_df = _with_extraction_timestamp(sample_df, extraction_timestamp)
    client_df = _with_extraction_timestamp(client_df, extraction_timestamp)
    classification_audit_df = _with_extraction_timestamp(classification_audit_df, extraction_timestamp)
    table_extraction_audit_df = _with_extraction_timestamp(table_extraction_audit_df, extraction_timestamp)
    section_extraction_audit_df = _with_extraction_timestamp(section_extraction_audit_df, extraction_timestamp)
    field_extraction_audit_df = _with_extraction_timestamp(field_extraction_audit_df, extraction_timestamp)
    duplicate_audit_df = _with_extraction_timestamp(duplicate_audit_df, extraction_timestamp)
    packaging_preservatives_df = _with_extraction_timestamp(packaging_preservatives_df, extraction_timestamp)
    notes_df = _with_extraction_timestamp(notes_df, extraction_timestamp)
    general_considerations_df = _with_extraction_timestamp(general_considerations_df, extraction_timestamp)
    conformity_statement_df = _with_extraction_timestamp(conformity_statement_df, extraction_timestamp)
    validation_key_df = _with_extraction_timestamp(validation_key_df, extraction_timestamp)
    revision_reason_df = _with_extraction_timestamp(revision_reason_df, extraction_timestamp)

    validations = validate_outputs(
        df,
        sample_df,
        client_df,
        packaging_preservatives_df,
        sheets_to_write,
    )
    validation_frames = [errors for errors in validations.values() if errors is not None and not errors.empty]
    validation_errors = (
        pd.concat(validation_frames, ignore_index=True)
        if validation_frames
        else pd.DataFrame(columns=VALIDATION_ERROR_COLUMNS)
    )
    validation_errors = _with_required_extraction_timestamp(validation_errors, extraction_timestamp)
    json_output_path = output_path.with_suffix(".jsonl")

    if ext == ".csv":
        df.to_csv(output_path, index=False, encoding="utf-8-sig")
        if not validation_errors.empty:
            validation_errors.to_csv(
                output_path.with_name(f"{output_path.stem}_validation_errors.csv"),
                index=False,
                encoding="utf-8-sig",
            )
    elif ext == ".json":
        df.to_json(output_path, orient="records", force_ascii=False, indent=2)
        if not validation_errors.empty:
            validation_errors.to_json(
                output_path.with_name(f"{output_path.stem}_validation_errors.json"),
                orient="records",
                force_ascii=False,
                indent=2,
            )
    else:
        frames = {
            "results_extract": df, "sample": sample_df, "client": client_df,
            "packaging_preservatives": packaging_preservatives_df, "notes": notes_df,
            "general_considerations": general_considerations_df,
            "conformity_statement": conformity_statement_df, "validation_key": validation_key_df,
            "revision_reason": revision_reason_df, "classification_audit": classification_audit_df,
            "table_extraction_audit": table_extraction_audit_df,
            "section_extraction_audit": section_extraction_audit_df,
            "field_extraction_audit": field_extraction_audit_df,
            "duplicate_audit": duplicate_audit_df, "validation_errors": validation_errors,
        }
        write_excel(
            output_path, sheets_to_write, frames,
            _format_numeric_results_sheet, _format_output_sheet,
        )

        _write_jsonl_output(
            json_output_path,
            df,
            sample_df,
            client_df,
            classification_audit_df,
            table_extraction_audit_df,
            section_extraction_audit_df,
            field_extraction_audit_df,
            duplicate_audit_df,
            packaging_preservatives_df,
            notes_df,
            general_considerations_df,
            conformity_statement_df,
            validation_key_df,
            revision_reason_df,
            validation_errors,
            sheets_to_write,
        )
    return validation_errors
