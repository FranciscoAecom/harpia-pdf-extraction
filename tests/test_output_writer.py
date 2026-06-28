import tempfile
import unittest
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from harpia_parser.constants import (
    CLIENT_COLUMNS,
    DUPLICATE_AUDIT_COLUMNS,
    FIELD_EXTRACTION_AUDIT_COLUMNS,
    PACKAGING_PRESERVATIVES_COLUMNS,
    RESULTS_EXTRACT_COLUMNS,
    SAMPLE_COLUMNS,
    SECTION_EXTRACTION_AUDIT_COLUMNS,
)
from harpia_parser.formatting.output_writer import salvar
from test_validation import _valid_row


class OutputWriterTest(unittest.TestCase):
    def test_acm_resultado_tratado_is_written_as_number_with_original_decimal_format(self):
        row = _valid_row()
        row["resultado"] = "< 0,100 mg/L"
        row["acm_resultado_tratado"] = 0.1
        row["lq"] = "0,0200 mg/L"
        row["acm_lq_minimo"] = 0.02
        row["acm_lq_maximo"] = 0.02
        row["incerteza"] = "0,030 %"
        row["acm_incerteza_valor"] = 0.03
        row["faixa_aceitacao"] = "75,00 a 125,000 %"
        row["acm_faixa_aceitacao_minimo"] = 75.0
        row["acm_faixa_aceitacao_maximo"] = 125.0
        row["variacao_percentual"] = "4,50"
        row["quantidade_adicionada"] = "10,000"
        row["recuperacao_percentual"] = "98,5"
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(df, output_path)

            workbook = load_workbook(output_path, data_only=False)
            worksheet = workbook["results_extract"]
            expected = {
                "acm_resultado_tratado": (0.1, "0.000"),
                "acm_lq_minimo": (0.02, "0.0000"),
                "acm_lq_maximo": (0.02, "0.0000"),
                "acm_incerteza_valor": (0.03, "0.000"),
                "acm_faixa_aceitacao_minimo": (75, "0.00"),
                "acm_faixa_aceitacao_maximo": (125, "0.000"),
                "variacao_percentual": (4.5, "0.00"),
                "quantidade_adicionada": (10, "0.000"),
                "recuperacao_percentual": (98.5, "0.0"),
            }
            for column_name, (value, number_format) in expected.items():
                column_index = RESULTS_EXTRACT_COLUMNS.index(column_name) + 1
                cell = worksheet.cell(row=2, column=column_index)
                self.assertEqual(cell.value, value)
                self.assertEqual(cell.number_format, number_format)

    def test_scientific_acm_resultado_tratado_is_written_as_number(self):
        row = _valid_row()
        row["resultado"] = "7,9 x 102"
        row["acm_resultado_tratado"] = 790
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(df, output_path)

            workbook = load_workbook(output_path, data_only=False)
            worksheet = workbook["results_extract"]
            column_index = RESULTS_EXTRACT_COLUMNS.index("acm_resultado_tratado") + 1
            cell = worksheet.cell(row=2, column=column_index)

            self.assertEqual(cell.value, 790)
            self.assertEqual(cell.number_format, "0.0")

    def test_common_ids_are_written_as_integers_and_sample_dates_as_dates(self):
        row = _valid_row()
        row["id_taxonomia"] = "1"
        row["versao_template"] = "2"
        row["id_amostra"] = "717727"
        row["acm_data_inicio"] = "04/12/2024 00:00:00"
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)
        sample_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": "1",
            "nome_taxonomia": "Agua Superficial",
            "versao_template": "2",
            "id_amostra": "717727",
            "data_coleta": "01/02/2025 08:30",
            "data_publicacao": "03/02/2025",
            "data_recebimento": "02/02/2025",
        }], columns=SAMPLE_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(df, output_path, sample_df=sample_df, output_tabs=["results_extract", "sample", "validation_errors"])

            workbook = load_workbook(output_path, data_only=False)
            results = workbook["results_extract"]
            sample = workbook["sample"]

            self.assertEqual(results.cell(row=2, column=RESULTS_EXTRACT_COLUMNS.index("id_taxonomia") + 1).value, 1)
            self.assertEqual(results.cell(row=2, column=RESULTS_EXTRACT_COLUMNS.index("versao_template") + 1).value, 2)
            self.assertEqual(results.cell(row=2, column=RESULTS_EXTRACT_COLUMNS.index("id_amostra") + 1).value, 717727)
            result_timestamp = results.cell(row=2, column=RESULTS_EXTRACT_COLUMNS.index("acm_data_hora_extracao") + 1)
            sample_timestamp = sample.cell(row=2, column=SAMPLE_COLUMNS.index("acm_data_hora_extracao") + 1)
            self.assertIsInstance(result_timestamp.value, datetime)
            self.assertEqual(sample_timestamp.value, result_timestamp.value)
            result_start_date = results.cell(row=2, column=RESULTS_EXTRACT_COLUMNS.index("acm_data_inicio") + 1)
            self.assertEqual(result_start_date.value, datetime(2024, 12, 4, 0, 0))
            self.assertEqual(result_start_date.number_format, "dd/mm/yyyy hh:mm")
            date_cell = sample.cell(row=2, column=SAMPLE_COLUMNS.index("data_coleta") + 1)
            self.assertEqual(date_cell.value, datetime(2025, 2, 1, 8, 30))
            self.assertEqual(date_cell.number_format, "dd/mm/yyyy hh:mm")
            publication_cell = sample.cell(row=2, column=SAMPLE_COLUMNS.index("data_publicacao") + 1)
            self.assertEqual(publication_cell.value, datetime(2025, 2, 3, 0, 0))
            self.assertEqual(publication_cell.number_format, "dd/mm/yyyy hh:mm")

    def test_table_extraction_audit_sheet_is_written_when_requested(self):
        row = _valid_row()
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)
        table_audit_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "pagina": 1,
            "tabela_indice": 1,
            "status": "ok",
        }])

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(
                df,
                output_path,
                output_tabs=["results_extract", "table_extraction_audit", "validation_errors"],
                table_extraction_audit_df=table_audit_df,
            )

            workbook = load_workbook(output_path, data_only=False)
            self.assertIn("table_extraction_audit", workbook.sheetnames)
            worksheet = workbook["table_extraction_audit"]
            self.assertEqual(worksheet.cell(row=2, column=1).value, "a.pdf")
            self.assertEqual(worksheet.freeze_panes, "A2")
            self.assertEqual(worksheet.cell(row=1, column=1).fill.fgColor.rgb, "001F4E78")
            self.assertEqual(worksheet.cell(row=1, column=1).font.color.rgb, "00FFFFFF")

    def test_duplicate_audit_sheet_is_written_when_requested(self):
        row = _valid_row()
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)
        duplicate_audit_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "caminho_arquivo": "C:/tmp/a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "status": "unico",
        }]).reindex(columns=DUPLICATE_AUDIT_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(
                df,
                output_path,
                output_tabs=["results_extract", "duplicate_audit", "validation_errors"],
                duplicate_audit_df=duplicate_audit_df,
            )

            workbook = load_workbook(output_path, data_only=False)
            self.assertIn("duplicate_audit", workbook.sheetnames)
            worksheet = workbook["duplicate_audit"]
            headers = [worksheet.cell(row=1, column=column).value for column in range(1, worksheet.max_column + 1)]
            payload = {
                str(header): worksheet.cell(row=2, column=column).value
                for column, header in enumerate(headers, start=1)
            }
            self.assertEqual(worksheet.cell(row=2, column=1).value, "a.pdf")
            self.assertEqual(payload["status"], "unico")
            self.assertIsInstance(payload["acm_data_hora_extracao"], datetime)

    def test_packaging_preservatives_sheet_is_written_when_requested(self):
        row = _valid_row()
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)
        packaging_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "id_amostra": "687944",
            "identificacao_amostra": "687944 - ECR 01R - P50",
            "embalagem": "Polietileno",
            "volume": "1000 mL",
            "preservacao": "0 a 6ºC",
            "metodos": "Sólidos Dissolvidos Totais.",
        }], columns=PACKAGING_PRESERVATIVES_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(
                df,
                output_path,
                output_tabs=["results_extract", "packaging_preservatives", "validation_errors"],
                packaging_preservatives_df=packaging_df,
            )

            workbook = load_workbook(output_path, data_only=False)
            self.assertIn("packaging_preservatives", workbook.sheetnames)
            worksheet = workbook["packaging_preservatives"]
            headers = [worksheet.cell(row=1, column=column).value for column in range(1, worksheet.max_column + 1)]
            payload = {
                str(header): worksheet.cell(row=2, column=column).value
                for column, header in enumerate(headers, start=1)
            }
            self.assertEqual(worksheet.cell(row=2, column=1).value, "a.pdf")
            self.assertEqual(payload["embalagem"], "Polietileno")
            self.assertIsInstance(payload["acm_data_hora_extracao"], datetime)
            self.assertEqual(worksheet.freeze_panes, "A2")

    def test_packaging_preservatives_validation_errors_are_written(self):
        row = _valid_row()
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)
        packaging_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "id_amostra": "687944",
            "identificacao_amostra": "687944 - ECR 01R - P50",
            "embalagem": "Polietileno",
            "volume": "1000 mL",
            "preservacao": "0 a 6ºC",
            "metodos": "",
        }], columns=PACKAGING_PRESERVATIVES_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(
                df,
                output_path,
                output_tabs=["results_extract", "packaging_preservatives", "validation_errors"],
                packaging_preservatives_df=packaging_df,
            )

            workbook = load_workbook(output_path, data_only=False)
            worksheet = workbook["validation_errors"]
            headers = [worksheet.cell(row=1, column=column).value for column in range(1, worksheet.max_column + 1)]
            row_values = [worksheet.cell(row=2, column=column).value for column in range(1, worksheet.max_column + 1)]
            payload = {
                str(header): row_values[index]
                for index, header in enumerate(headers)
                if index < len(row_values)
            }

            self.assertEqual(payload["sheet"], "packaging_preservatives")
            self.assertEqual(payload["field"], "metodos")
            self.assertIsInstance(payload["acm_data_hora_extracao"], datetime)

    def test_sample_and_client_validation_errors_are_written(self):
        row = _valid_row()
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)
        sample_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "id_amostra": "687944",
            "identificacao_amostra": "687944 - ECR 01R - P50",
            "tipo_amostra": "Agua superficial",
            "criterio_conformidade": None,
            "data_coleta": "2025-01-22",
            "data_publicacao": "22/01/2025",
            "data_recebimento": "22/01/2025",
            "observacoes": None,
            "localizacao": None,
            "latitude": "-19,123",
            "longitude": "-43.123",
            "condicoes_climaticas_nas_ultimas_24_horas": None,
            "condicoes_climaticas_no_momento_da_coleta": None,
            "tipo_coleta": None,
            "responsavel_amostra": None,
            "planejamento_amostragem": None,
            "descricao_nao_conformidade": None,
        }], columns=SAMPLE_COLUMNS)
        client_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "id_amostra": "",
            "proposta_comercial": "PC-1",
            "cliente": "Cliente",
            "cnpj_cpf": "00.000.000/0001-00",
            "contato": None,
            "telefone": None,
            "endereco": None,
        }], columns=CLIENT_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(
                df,
                output_path,
                sample_df=sample_df,
                client_df=client_df,
                output_tabs=["results_extract", "sample", "client", "validation_errors"],
            )

            workbook = load_workbook(output_path, data_only=False)
            worksheet = workbook["validation_errors"]
            headers = [worksheet.cell(row=1, column=column).value for column in range(1, worksheet.max_column + 1)]
            rows = [
                {
                    str(header): worksheet.cell(row=row, column=column).value
                    for column, header in enumerate(headers, start=1)
                }
                for row in range(2, worksheet.max_row + 1)
            ]
            observed = {(row["sheet"], row["field"]) for row in rows}

            self.assertIn(("sample", "data_coleta"), observed)
            self.assertIn(("client", "id_amostra"), observed)

    def test_document_section_sheets_are_written_when_requested(self):
        row = _valid_row()
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)
        notes_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "id_amostra": "687944",
            "pagina": 5,
            "texto": "Texto de nota.",
        }])
        general_considerations_df = notes_df.copy()
        conformity_statement_df = notes_df.copy()
        validation_key_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "id_amostra": "687944",
            "pagina": 5,
            "chave_validacao": "ABC-123",
        }])

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(
                df,
                output_path,
                output_tabs=[
                    "results_extract",
                    "notes",
                    "general_considerations",
                    "conformity_statement",
                    "validation_key",
                    "validation_errors",
                ],
                notes_df=notes_df,
                general_considerations_df=general_considerations_df,
                conformity_statement_df=conformity_statement_df,
                validation_key_df=validation_key_df,
            )

            workbook = load_workbook(output_path, data_only=False)
            for sheet_name in ["notes", "general_considerations", "conformity_statement", "validation_key"]:
                self.assertIn(sheet_name, workbook.sheetnames)
            self.assertEqual(workbook["validation_key"].cell(row=2, column=6).value, "ABC-123")

    def test_json_output_is_written_grouped_by_document(self):
        row = _valid_row()
        row["nome_do_arquivo"] = "a.pdf"
        row["id_taxonomia"] = 1
        row["nome_taxonomia"] = "Agua Superficial"
        row["versao_template"] = "1"
        row["id_amostra"] = "687944"
        row["acm_data_inicio"] = "04/12/2024 00:00:00"
        row["lq"] = "2,00 - 12,000"
        row["acm_lq_minimo"] = 2.0
        row["acm_lq_maximo"] = 12.0
        row["variacao_percentual"] = "4,50"
        row["quantidade_adicionada"] = "10,000"
        row["recuperacao_percentual"] = "98,5"
        df = pd.DataFrame([row], columns=RESULTS_EXTRACT_COLUMNS)
        sample_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "versao_template": "1",
            "id_amostra": "687944",
            "identificacao_amostra": "687944 - ECR 01R - P50",
            "tipo_amostra": "Agua superficial",
            "data_coleta": "01/02/2025 08:30",
            "data_publicacao": "03/02/2025",
        }])
        table_audit_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "versao_template": "1",
            "status": "ok",
        }])

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "extracted_data.xlsx"
            salvar(
                df,
                output_path,
                sample_df=sample_df,
                output_tabs=[
                    "results_extract",
                    "sample",
                    "table_extraction_audit",
                    "duplicate_audit",
                    "validation_errors",
                ],
                table_extraction_audit_df=table_audit_df,
                duplicate_audit_df=pd.DataFrame([{
                    "nome_do_arquivo": "a.pdf",
                    "status": "unico",
                }]).reindex(columns=DUPLICATE_AUDIT_COLUMNS),
            )

            json_path = output_path.with_suffix(".jsonl")
            self.assertTrue(json_path.exists())
            json_text = json_path.read_text(encoding="utf-8")
            document = json.loads(json_text.splitlines()[0])

            self.assertEqual(document["arquivo"]["nome_do_arquivo"], "a.pdf")
            self.assertIn('"acm_resultado_tratado":7.100', json_text)
            self.assertIn('"acm_lq_minimo":2.00', json_text)
            self.assertIn('"acm_lq_maximo":12.000', json_text)
            self.assertIn('"variacao_percentual":4.50', json_text)
            self.assertIn('"quantidade_adicionada":10.000', json_text)
            self.assertIn('"recuperacao_percentual":98.5', json_text)
            self.assertEqual(document["arquivo"]["id_taxonomia"], 1)
            self.assertIn("acm_data_hora_extracao", document["tabelas"]["results_extract"][0])
            self.assertEqual(document["tabelas"]["results_extract"][0]["id_amostra"], 687944)
            self.assertEqual(document["tabelas"]["results_extract"][0]["acm_data_inicio"], "04/12/2024 00:00:00")
            self.assertEqual(document["tabelas"]["sample"][0]["identificacao_amostra"], "687944 - ECR 01R - P50")
            self.assertEqual(document["tabelas"]["sample"][0]["data_coleta"], "01/02/2025 08:30:00")
            self.assertEqual(document["tabelas"]["sample"][0]["data_publicacao"], "03/02/2025 00:00:00")
            self.assertNotIn("auditoria", document)


    def test_section_extraction_audit_is_written_only_to_excel(self):
        df = pd.DataFrame([_valid_row()], columns=RESULTS_EXTRACT_COLUMNS)
        section_audit_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "versao_template": 1,
            "pagina": 5,
            "objeto_tipo": "secao",
            "secao": "notes",
            "titulo_detectado": "Notas",
            "modo_auditoria": "conhecida",
            "ocorrencias_detectadas": 1,
            "registros_extraidos": 1,
            "status": "ok",
            "observacao": "Secao reconhecida e dados extraidos.",
        }]).reindex(columns=SECTION_EXTRACTION_AUDIT_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(
                df,
                output_path,
                output_tabs=["results_extract", "section_extraction_audit", "validation_errors"],
                section_extraction_audit_df=section_audit_df,
            )

            workbook = load_workbook(output_path, data_only=True)
            self.assertIn("section_extraction_audit", workbook.sheetnames)
            self.assertEqual(workbook["section_extraction_audit"].cell(row=2, column=15).value, "ok")

            document = json.loads(output_path.with_suffix(".jsonl").read_text(encoding="utf-8").splitlines()[0])
            self.assertNotIn("auditoria", document)

    def test_field_extraction_audit_is_written_only_to_excel(self):
        df = pd.DataFrame([_valid_row()], columns=RESULTS_EXTRACT_COLUMNS)
        field_audit_df = pd.DataFrame([{
            "nome_do_arquivo": "a.pdf",
            "id_taxonomia": 1,
            "nome_taxonomia": "Agua Superficial",
            "versao_template": 1,
            "schema_origem": "sample",
            "campo": "latitude_real",
            "pagina": 1,
            "regra_encontrada": "Latitude real",
            "ocorrencias_detectadas": 1,
            "valor_extraido": "-19,23234",
            "status": "ok",
        }]).reindex(columns=FIELD_EXTRACTION_AUDIT_COLUMNS)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "out.xlsx"
            salvar(
                df,
                output_path,
                output_tabs=["results_extract", "field_extraction_audit", "validation_errors"],
                field_extraction_audit_df=field_audit_df,
            )

            workbook = load_workbook(output_path, data_only=True)
            self.assertIn("field_extraction_audit", workbook.sheetnames)
            self.assertEqual(workbook["field_extraction_audit"].cell(row=2, column=13).value, "ok")

            document = json.loads(output_path.with_suffix(".jsonl").read_text(encoding="utf-8").splitlines()[0])
            self.assertNotIn("auditoria", document)


if __name__ == "__main__":
    unittest.main()

